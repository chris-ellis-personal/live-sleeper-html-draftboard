#!/usr/bin/env python3
"""Generates examples/sample_rankings.xlsx, a tiny rankings sheet in the layout
the tool expects. The player order is illustrative only -- not real advice.

    python examples/make_sample_rankings.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

# (name, team, bye, position) in overall-rank order. A few names are chosen on
# purpose because they exercise name matching: apostrophe, hyphen, "Jr." / "III".
PLAYERS = [
    ("Jahmyr Gibbs", "DET", 6, "RB"),
    ("Ja'Marr Chase", "CIN", 6, "WR"),
    ("Bijan Robinson", "ATL", 11, "RB"),
    ("Puka Nacua", "LAR", 11, "WR"),
    ("Amon-Ra St. Brown", "DET", 6, "WR"),
    ("Christian McCaffrey", "SF", 8, "RB"),
    ("Josh Allen", "BUF", 7, "QB"),
    ("Brock Bowers", "LV", 13, "TE"),
    ("CeeDee Lamb", "DAL", 14, "WR"),
    ("Saquon Barkley", "PHI", 10, "RB"),
    ("Lamar Jackson", "BAL", 13, "QB"),
    ("Justin Jefferson", "MIN", 6, "WR"),
    ("Jonathan Taylor", "IND", 13, "RB"),
    ("Trey McBride", "ARI", 14, "TE"),
    ("James Cook III", "BUF", 7, "RB"),
    ("Jaxon Smith-Njigba", "SEA", 11, "WR"),
    ("Joe Burrow", "CIN", 6, "QB"),
    ("Derrick Henry", "BAL", 13, "RB"),
    ("A.J. Brown", "NE", 11, "WR"),
    ("Jayden Daniels", "WAS", 7, "QB"),
    ("Tyler Warren", "IND", 13, "TE"),
    ("Jalen Hurts", "PHI", 10, "QB"),
    ("Sam LaPorta", "DET", 6, "TE"),
    ("Drake Maye", "NE", 11, "QB"),
    ("Tucker Kraft", "GB", 11, "TE"),
    ("Harold Fannin Jr.", "CLE", 11, "TE"),
]

# Order of the optional per-position "tier blocks": (position, first column index)
BLOCKS = [("QB", 6), ("WR", 13), ("RB", 20), ("TE", 27)]  # F, M, T, AA


def build(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rankings"
    bold = Font(bold=True)

    for col, text in enumerate(["PLAYER NAME", "TEAM", "BYE", "POS."], start=1):
        ws.cell(row=1, column=col, value=text).font = bold
    for i, (name, team, bye, pos) in enumerate(PLAYERS, start=2):
        for col, val in enumerate((name, team, bye, pos), start=1):
            ws.cell(row=i, column=col, value=val)

    # Tier blocks: "<POS> TIERS" header, then name/team/bye/pos/overall rank/tier.
    # In a real sheet these are usually =FILTER() formulas; static values work too.
    for pos, start in BLOCKS:
        ws.cell(row=1, column=start, value=f"{pos} TIERS").font = bold
        ws.cell(row=1, column=start + 4, value="Overall rank").font = bold
        ws.cell(row=1, column=start + 5, value="Tier rank").font = bold
        rows = [(rank, p) for rank, p in enumerate(PLAYERS, start=1) if p[3] == pos]
        for n, (overall, (name, team, bye, _)) in enumerate(rows, start=2):
            for off, val in enumerate((name, team, bye, pos, overall, 1 + (n - 2) // 4)):
                ws.cell(row=n, column=start + off, value=val)

    ws.column_dimensions["A"].width = 24
    wb.save(path)


if __name__ == "__main__":
    out = Path(__file__).with_name("sample_rankings.xlsx")
    build(out)
    print(f"Wrote {out}")
