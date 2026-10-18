"""Command-line interface: `python -m sleeper_draftboard run|revert ...`"""

import argparse
import shutil
import sys
import traceback
import webbrowser
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("This tool requires 'openpyxl'. Install it with:\n    pip install openpyxl")

from . import __version__
from .config import POLL_INTERVAL_SECONDS, default_state_dir
from .rankings import LayoutError, Rankings
from .sleeper import parse_draft_id
from .state import acquire_lock, load_state, new_state, release_lock, save_state, state_paths
from .tracker import DraftTracker


def _clean_path(p):
    """Strip surrounding quotes/whitespace (Windows 'Copy as path' adds quotes)."""
    return p if p is None else p.strip().strip('"').strip("'")


def _fail(msg):
    print(f"Error: {msg}")
    sys.exit(1)


def _get_draft_id(arg_value, prompt):
    raw = _clean_path(arg_value or input(prompt))
    try:
        return parse_draft_id(raw)
    except ValueError as e:
        _fail(str(e))


def _pick_sheet(wb, sheet):
    if not sheet:
        return wb.active
    if sheet in wb.sheetnames:
        return wb[sheet]
    if sheet.isdigit() and 0 <= int(sheet) < len(wb.sheetnames):
        return wb[wb.sheetnames[int(sheet)]]
    _fail(f"sheet '{sheet}' not found. Available sheets: {', '.join(wb.sheetnames)}")


def cmd_run(args):
    if not args.file:
        args.file = _clean_path(input("Path to your rankings .xlsx: "))
    source = Path(_clean_path(args.file)).expanduser()
    if not source.exists():
        _fail(f"path not found: {source}")
    if source.is_dir():
        _fail(f"'{source}' is a folder, not a file. Point --file at the .xlsx itself.")
    if source.suffix.lower() not in (".xlsx", ".xlsm"):
        print(f"Warning: '{source}' doesn't look like an .xlsx/.xlsm file. Continuing anyway.")

    draft_id = _get_draft_id(args.draft, "Paste the Sleeper draft URL or ID: ")

    state_dir = Path(args.state_dir).expanduser() if args.state_dir else default_state_dir()
    paths = state_paths(state_dir, draft_id)
    state = load_state(paths["state"])

    out_dir = Path(args.output_dir).expanduser() if args.output_dir else source.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    html_path = out_dir / f"{source.stem}_live_view.html"

    # Validate the layout on the pristine source before creating anything.
    try:
        probe = openpyxl.load_workbook(source)
        Rankings(_pick_sheet(probe, args.sheet), header_row=args.header_row)
    except LayoutError as e:
        _fail(str(e))

    working_copy = None
    if not args.no_excel:
        working_copy = state.get("working_copy")
        if not working_copy or not Path(working_copy).exists():
            working_copy = str(out_dir / f"{source.stem}_LIVE{source.suffix}")
            shutil.copy2(source, working_copy)
            fresh = new_state()
            state.update({k: fresh[k] for k in (
                "processed_pick_nos", "struck_coords", "drafted_master_rows",
                "unmatched_entries")})
            state["working_copy"] = working_copy
            save_state(paths["state"], state)
            print(f"Created working copy: {working_copy}")
        else:
            print(f"Resuming existing working copy: {working_copy}")

    try:
        acquire_lock(paths["lock"])
    except RuntimeError as e:
        _fail(str(e))

    try:
        wb = openpyxl.load_workbook(working_copy or source)
        ws = _pick_sheet(wb, args.sheet)
        rankings = Rankings(ws, header_row=args.header_row)
        poll = args.poll_interval
        tracker = DraftTracker(
            draft_id, rankings, paths, state, html_path,
            wb=None if args.no_excel else wb,
            working_copy=working_copy,
            whole_row=args.whole_row,
            refresh=args.refresh or poll,
            poll_interval=poll,
            title=args.title,
        )
        opener = (lambda: webbrowser.open(html_path.resolve().as_uri())) \
            if args.open_browser else None
        tracker.run(open_browser_fn=opener)
    except KeyboardInterrupt:
        print("\nStopped by user. Progress has been saved; re-run 'run' to resume.")
    finally:
        release_lock(paths["lock"])


def cmd_revert(args):
    draft_id = _get_draft_id(args.draft, "Paste the Sleeper draft URL or ID to revert: ")
    state_dir = Path(args.state_dir).expanduser() if args.state_dir else default_state_dir()
    paths = state_paths(state_dir, draft_id)
    state = load_state(paths["state"])
    working_copy = state.get("working_copy")

    # Only ever delete files this tool created (they end in _LIVE).
    if working_copy and Path(working_copy).exists() and Path(working_copy).stem.endswith("_LIVE"):
        Path(working_copy).unlink()
        print(f"Deleted working copy: {working_copy}")
    else:
        print("No working copy found for this draft (nothing to revert).")

    for p in (paths["state"], paths["lock"], paths["unmatched_log"]):
        if p.exists():
            p.unlink()

    print("Note: the HTML view file (if any) was left in place -- delete it manually if you want it gone.")
    print("Your original rankings file was never touched.")


def build_arg_parser():
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--state-dir",
                        help="Where progress/lock/alias files are kept "
                             "(default: ~/.sleeper_draftboard, or $SLEEPER_DRAFTBOARD_STATE_DIR).")

    parser = argparse.ArgumentParser(
        prog="sleeper-draftboard",
        description="Live, auto-refreshing HTML draft board for a Sleeper draft, "
                    "driven by your own rankings spreadsheet.")
    parser.add_argument("--version", action="version", version=f"%(prog)s {__version__}")
    sub = parser.add_subparsers(dest="command")

    p_run = sub.add_parser("run", parents=[common],
                           help="Start/resume tracking a draft and keep the HTML view updated.")
    p_run.add_argument("--file", help="Path to your rankings .xlsx.")
    p_run.add_argument("--draft", help="Sleeper draft URL or ID.")
    p_run.add_argument("--sheet", help="Worksheet name or 0-based index (default: the active sheet).")
    p_run.add_argument("--header-row", type=int, default=1,
                       help="Row containing your column headers (default: %(default)s).")
    p_run.add_argument("--output-dir",
                       help="Where to write the HTML view and Excel working copy "
                            "(default: the folder containing your rankings file).")
    p_run.add_argument("--title", default="Live Draft Board", help="Heading shown on the board.")
    p_run.add_argument("--whole-row", action="store_true",
                       help="In the Excel copy, strike the whole row instead of just the name.")
    p_run.add_argument("--no-excel", action="store_true",
                       help="Don't create/update an .xlsx working copy; HTML view only.")
    p_run.add_argument("--poll-interval", type=int, default=POLL_INTERVAL_SECONDS,
                       help="Seconds between Sleeper API checks (default: %(default)s).")
    p_run.add_argument("--refresh", type=int,
                       help="Browser auto-refresh interval in seconds (default: the poll interval).")
    p_run.add_argument("--open-browser", action="store_true",
                       help="Open the HTML view in your default browser on start.")
    p_run.set_defaults(func=cmd_run)

    p_revert = sub.add_parser("revert", parents=[common],
                              help="Discard the working copy and saved state for a draft.")
    p_revert.add_argument("--draft", help="Sleeper draft URL or ID.")
    p_revert.set_defaults(func=cmd_revert)
    return parser


def interactive_menu():
    print("Sleeper Draft -- Live HTML Board")
    print("=" * 40)
    print("1) run    - start/resume tracking a draft")
    print("2) revert - discard a draft's working copy")
    choice = input("Choose an option [1/2]: ").strip()

    parser = build_arg_parser()
    if choice == "2":
        return parser.parse_args(["revert"])
    argv = ["run"]
    if input("Skip the Excel working copy and use HTML only? [y/N]: ").strip().lower() == "y":
        argv.append("--no-excel")
    if input("Open the board in your browser automatically? [Y/n]: ").strip().lower() != "n":
        argv.append("--open-browser")
    return parser.parse_args(argv)  # --file / --draft are prompted for inside cmd_run


def main(argv=None):
    parser = build_arg_parser()
    args = parser.parse_args(argv)
    if not args.command:
        args = interactive_menu()
    try:
        args.func(args)
    except KeyboardInterrupt:
        print("\nInterrupted.")
    except Exception:
        print("\nUnexpected error:")
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    main()
