"""Reads a rankings spreadsheet and works out its layout from the header row.

Required: a header row containing a player-name column and a position column
(a team column and a bye column are optional). Optional: "tier blocks", i.e.
per-position panels headed e.g. "QB TIERS" whose rows mirror the master list
filtered by position (typically built with a FILTER() formula). When blocks
are present, the Excel working copy gets strikethrough applied in them too.
"""

import re
from copy import copy
from dataclasses import dataclass, field
from typing import Dict, List, Optional

from openpyxl.utils import get_column_letter

from .config import PREFERRED_POSITION_ORDER
from .names import normalize_name

HEADER_ALIASES = {
    "name": {"player name", "player", "name"},
    "team": {"team", "tm"},
    "bye": {"bye", "bye week"},
    "pos": {"pos.", "pos", "position"},
}

BLOCK_HEADER_RE = re.compile(r"^\s*([A-Za-z]{1,3})\s+TIERS?\s*$", re.IGNORECASE)

# Column offsets inside a tier block, counted from the "<POS> TIERS" header cell.
BLOCK_OFFSETS = {"name": 0, "team": 1, "bye": 2, "pos": 3, "overall": 4, "tier": 5}

# Rows to keep scanning past a blank name cell before deciding the list is over.
MAX_BLANK_GAP = 5


class LayoutError(Exception):
    pass


@dataclass
class Player:
    name: str
    team: str
    bye: object
    position: str
    master_row: int
    tier_row: Optional[int] = None
    tier: Optional[int] = None


@dataclass
class SheetLayout:
    header_row: int
    name_col: int
    pos_col: int
    team_col: Optional[int] = None
    bye_col: Optional[int] = None
    blocks: Dict[str, Dict[str, int]] = field(default_factory=dict)  # pos -> field -> col index

    @property
    def first_data_row(self):
        return self.header_row + 1

    def describe(self):
        def col(c):
            return get_column_letter(c) if c else "-"
        parts = [f"name={col(self.name_col)} team={col(self.team_col)} "
                 f"bye={col(self.bye_col)} pos={col(self.pos_col)}"]
        if self.blocks:
            parts.append("tier blocks: " + ", ".join(
                f"{p}@{col(b['name'])}" for p, b in self.blocks.items()))
        else:
            parts.append("tier blocks: none")
        return "; ".join(parts)


def detect_layout(ws, header_row=1):
    """Finds the master-list columns and any tier blocks from the header row."""
    found = {}
    blocks = {}
    for cell in ws[header_row]:
        if cell.value is None:
            continue
        text = str(cell.value).strip()
        low = text.lower()
        m = BLOCK_HEADER_RE.match(text)
        if m:
            start = cell.column
            blocks[m.group(1).upper()] = {k: start + off for k, off in BLOCK_OFFSETS.items()}
            continue
        for key, aliases in HEADER_ALIASES.items():
            if low in aliases and key not in found:
                found[key] = cell.column
                break

    missing = [k for k in ("name", "pos") if k not in found]
    if missing:
        labels = {"name": "a player-name column (e.g. 'PLAYER NAME')",
                  "pos": "a position column (e.g. 'POS.')"}
        raise LayoutError(
            f"Could not find {' and '.join(labels[k] for k in missing)} in row "
            f"{header_row} of sheet '{ws.title}'. If your headers are on a different "
            f"row, pass --header-row N. See the README for the expected format."
        )

    return SheetLayout(
        header_row=header_row,
        name_col=found["name"],
        pos_col=found["pos"],
        team_col=found.get("team"),
        bye_col=found.get("bye"),
        blocks=blocks,
    )


class Rankings:
    """In-memory index of the master list."""

    def __init__(self, ws, layout=None, header_row=1):
        self.ws = ws
        self.layout = layout or detect_layout(ws, header_row)
        self.players: List[Player] = []
        self.by_norm_name: Dict[str, List[Player]] = {}
        self.positions: List[str] = []
        self.max_row = self.layout.first_data_row - 1
        self._build()

    def _cell(self, row, col):
        return self.ws.cell(row=row, column=col).value if col else None

    def _build(self):
        lay = self.layout
        counters = {pos: 0 for pos in lay.blocks}
        blank_run = 0
        row = lay.first_data_row
        seen_positions = []

        while blank_run < MAX_BLANK_GAP and row <= max(self.ws.max_row, lay.first_data_row):
            name = self._cell(row, lay.name_col)
            if name is None or str(name).strip() == "":
                blank_run += 1
                row += 1
                continue
            blank_run = 0

            pos = str(self._cell(row, lay.pos_col) or "").strip().upper()
            tier_row = tier = None
            if pos in lay.blocks:
                counters[pos] += 1
                tier_row = lay.header_row + counters[pos]
                raw = self._cell(tier_row, lay.blocks[pos]["tier"])
                try:
                    tier = int(raw) if raw is not None else None
                except (TypeError, ValueError):
                    tier = None

            team = self._cell(row, lay.team_col)
            player = Player(
                name=str(name).strip(),
                team="" if team is None else str(team).strip(),
                bye=self._cell(row, lay.bye_col),
                position=pos,
                master_row=row,
                tier_row=tier_row,
                tier=tier,
            )
            self.players.append(player)
            self.by_norm_name.setdefault(normalize_name(player.name), []).append(player)
            if pos and pos not in seen_positions:
                seen_positions.append(pos)
            self.max_row = row
            row += 1

        preferred = [p for p in PREFERRED_POSITION_ORDER if p in seen_positions]
        self.positions = preferred + [p for p in seen_positions if p not in preferred]

    def lookup_exact(self, normalized_name):
        return self.by_norm_name.get(normalized_name)

    def lookup_by_position(self, position):
        return [p for p in self.players if p.position == position]


# --------------------------------------------------------------------------
# Excel strikethrough (only used for the optional .xlsx working copy)
# --------------------------------------------------------------------------

def set_strike(ws, row, col, strike=True):
    """Toggle strikethrough on a cell, preserving every other font attribute."""
    cell = ws.cell(row=row, column=col)
    font = copy(cell.font)
    font.strike = strike
    cell.font = font


def strike_player(ws, layout, player, strike=True, whole_row=False):
    """Apply/remove strikethrough for a player in the master list and, if the
    sheet has one for their position, the mirrored tier block. Returns the
    list of cell coordinates touched."""
    touched = []
    keys = ["name", "team", "bye", "pos"] if whole_row else ["name"]

    master_cols = {"name": layout.name_col, "team": layout.team_col,
                   "bye": layout.bye_col, "pos": layout.pos_col}
    for k in keys:
        col = master_cols[k]
        if col:
            set_strike(ws, player.master_row, col, strike)
            touched.append(f"{get_column_letter(col)}{player.master_row}")

    block = layout.blocks.get(player.position)
    if block and player.tier_row:
        for k in keys:
            col = block[k]
            set_strike(ws, player.tier_row, col, strike)
            touched.append(f"{get_column_letter(col)}{player.tier_row}")
    return touched
