import sys
from pathlib import Path

import openpyxl
import pytest

from sleeper_draftboard import cli, sleeper
from sleeper_draftboard.names import NameMatcher, normalize_name
from sleeper_draftboard.rankings import LayoutError, Rankings, detect_layout, strike_player
from sleeper_draftboard.state import state_paths
from sleeper_draftboard.tracker import DraftTracker

SAMPLE = Path(__file__).resolve().parent.parent / "examples" / "sample_rankings.xlsx"


def load_sample():
    wb = openpyxl.load_workbook(SAMPLE)
    return wb, Rankings(wb.active)


def pick(no, first, last, pos, team="XXX"):
    return {"pick_no": no, "metadata": {"first_name": first, "last_name": last,
                                         "position": pos, "team": team}}


def test_normalize_name():
    assert normalize_name("Ja'Marr Chase") == "jamarr chase"
    assert normalize_name("Harold Fannin Jr.") == "harold fannin"
    assert normalize_name("James Cook III") == "james cook"
    assert normalize_name("A.J. Brown") == "aj brown"
    assert normalize_name("") == ""


@pytest.mark.parametrize("raw", [
    "https://sleeper.com/draft/nfl/1399811152808034304?ftue=commish",
    "https://sleeper.com/draft/nfl/1399811152808034304",
    "1399811152808034304",
])
def test_parse_draft_id(raw):
    assert sleeper.parse_draft_id(raw) == "1399811152808034304"


def test_parse_draft_id_rejects_garbage():
    with pytest.raises(ValueError):
        sleeper.parse_draft_id("not a draft")


def test_layout_detection_on_sample():
    wb, r = load_sample()
    lay = r.layout
    assert (lay.name_col, lay.team_col, lay.bye_col, lay.pos_col) == (1, 2, 3, 4)
    assert set(lay.blocks) == {"QB", "WR", "RB", "TE"}
    assert len(r.players) == 26
    assert r.positions == ["QB", "RB", "WR", "TE"]


def test_layout_without_blocks_and_with_shuffled_columns():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Rank", "Position", "Player"])
    ws.append([1, "RB", "Some Guy"])
    r = Rankings(ws)
    assert r.layout.name_col == 3 and r.layout.pos_col == 2
    assert r.layout.blocks == {} and r.layout.team_col is None
    assert r.players[0].name == "Some Guy"


def test_missing_headers_raise_helpful_error():
    wb = openpyxl.Workbook()
    wb.active.append(["foo", "bar"])
    with pytest.raises(LayoutError) as e:
        detect_layout(wb.active)
    assert "--header-row" in str(e.value)


def test_header_row_option():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["My 2026 rankings"])
    ws.append(["PLAYER NAME", "TEAM", "BYE", "POS."])
    ws.append(["Josh Allen", "BUF", 7, "QB"])
    r = Rankings(ws, header_row=2)
    assert [p.name for p in r.players] == ["Josh Allen"]


def test_matcher_exact_fuzzy_and_alias(tmp_path):
    _, r = load_sample()
    m = NameMatcher(r, alias_path=tmp_path / "aliases.json")
    p, t = m.match("Ja'Marr Chase", "WR")
    assert p.name == "Ja'Marr Chase" and t == "exact"
    p, t = m.match("Harold Fannin", "TE")           # suffix dropped by Sleeper
    assert p.name == "Harold Fannin Jr." and t == "exact"
    p, t = m.match("Jaxon Smith Njigba", "WR")      # hyphen missing -> fuzzy
    assert p.name == "Jaxon Smith-Njigba" and t.startswith("fuzzy")
    p, t = m.match("Jaxon Smith Njigba", "WR")      # now remembered
    assert t == "alias"
    assert m.match("Totally Unknown Person", "WR") == (None, "none")


def test_optional_positions_do_not_fuzzy_match(tmp_path):
    _, r = load_sample()
    m = NameMatcher(r, alias_path=tmp_path / "a.json")
    assert m.match("Harrison Butker", "K") == (None, "none")


def test_tracker_marks_drafted_and_renders_html(tmp_path):
    wb, r = load_sample()
    paths = state_paths(tmp_path / "state", "123")
    from sleeper_draftboard.state import new_state
    html = tmp_path / "board.html"
    t = DraftTracker("123", r, paths, new_state(), html, wb=wb,
                     working_copy=str(tmp_path / "w.xlsx"), whole_row=False)
    all_picks = [
        pick(1, "Jahmyr", "Gibbs", "RB"),
        pick(2, "Ja'Marr", "Chase", "WR"),
        pick(3, "Some", "Rookie", "WR"),            # not in rankings
        pick(4, "Harrison", "Butker", "K"),         # silently skipped
        pick(5, "Dallas", "Cowboys", "DEF"),        # silently skipped
    ]
    changed = t.process_picks(all_picks)
    assert changed
    assert len(t.drafted_rows) == 2
    assert [u[1] for u in t.unmatched] == ["Some Rookie"]
    t.persist()
    t.render()
    page = html.read_text(encoding="utf-8")
    assert page.count('class="drafted"') == 4        # 2 players x (Overall + position col)
    assert "Some Rookie" in page and "Unmatched picks" in page

    # Excel copy: strikethrough applied in master list and tier block
    out = openpyxl.load_workbook(tmp_path / "w.xlsx").active
    gibbs = r.lookup_exact("jahmyr gibbs")[0]
    assert out.cell(gibbs.master_row, 1).font.strike is True
    assert out.cell(gibbs.tier_row, r.layout.blocks["RB"]["name"]).font.strike is True
    other = r.lookup_exact("josh allen")[0]
    assert not out.cell(other.master_row, 1).font.strike

    # Re-processing the same picks is a no-op
    assert t.process_picks(all_picks) is False


def test_html_only_mode_never_touches_workbook(tmp_path):
    wb, r = load_sample()
    from sleeper_draftboard.state import new_state
    t = DraftTracker("1", r, state_paths(tmp_path, "1"), new_state(), tmp_path / "b.html")
    t.process_picks([pick(1, "Josh", "Allen", "QB")])
    assert t.struck_coords == set() and len(t.drafted_rows) == 1


def test_draft_reset_detected(tmp_path):
    wb, r = load_sample()
    from sleeper_draftboard.state import new_state
    t = DraftTracker("1", r, state_paths(tmp_path, "1"), new_state(), tmp_path / "b.html")
    t.process_picks([pick(1, "Josh", "Allen", "QB"), pick(2, "Lamar", "Jackson", "QB")])
    t.process_picks([pick(1, "Joe", "Burrow", "QB")])
    assert len(t.drafted_rows) == 1


def test_is_complete():
    info = {"status": "drafting", "settings": {"rounds": 2, "teams": 2}}
    assert not DraftTracker.is_complete(info, [1, 2, 3])
    assert DraftTracker.is_complete(info, [1, 2, 3, 4])
    assert DraftTracker.is_complete({"status": "complete"}, [])


def test_end_to_end_cli_with_mocked_sleeper(tmp_path, monkeypatch, capsys):
    """Runs the real CLI through a full 'draft' using a fake API and a copy of
    the sample outside the repo, to prove nothing depends on folder layout."""
    work = tmp_path / "some" / "random folder"
    work.mkdir(parents=True)
    rankings = work / "my rankings.xlsx"
    rankings.write_bytes(SAMPLE.read_bytes())

    picks = [pick(1, "Jahmyr", "Gibbs", "RB"), pick(2, "Josh", "Allen", "QB")]
    monkeypatch.setattr(sleeper, "get_draft_info",
                        lambda d: {"status": "complete", "settings": {"rounds": 1, "teams": 2}})
    monkeypatch.setattr(sleeper, "get_picks", lambda d: picks)
    monkeypatch.chdir(tmp_path)  # unrelated cwd

    cli.main(["run", "--file", str(rankings),
              "--draft", "https://sleeper.com/draft/nfl/999999999999",
              "--state-dir", str(tmp_path / "st")])

    assert (work / "my rankings_live_view.html").exists()
    assert (work / "my rankings_LIVE.xlsx").exists()
    assert (tmp_path / "st" / "draft_999999999999.json").exists()
    assert not (tmp_path / ".sleeper_draftboard").exists()

    # revert removes only the generated working copy, never the source
    cli.main(["revert", "--draft", "999999999999", "--state-dir", str(tmp_path / "st")])
    assert not (work / "my rankings_LIVE.xlsx").exists()
    assert rankings.exists()


def test_cli_rejects_bad_layout(tmp_path, capsys):
    bad = tmp_path / "bad.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["nothing", "useful"])
    wb.save(bad)
    with pytest.raises(SystemExit):
        cli.main(["run", "--file", str(bad), "--draft", "1234567890",
                  "--state-dir", str(tmp_path / "st")])
    assert "Could not find" in capsys.readouterr().out
